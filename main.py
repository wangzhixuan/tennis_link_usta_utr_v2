import os
import argparse
import logging
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import USTA_USER, USTA_PASS, UTR_USER, UTR_PASS, HEADLESS
from usta_scraper import USTAScraper
from utr_scraper import UTRScraper
from matcher import PlayerMatcher

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

def parse_args():
    parser = argparse.ArgumentParser(
        description="TennisLink: Fetch USTA players, match them with UTR ratings, and export a beautiful spreadsheet."
    )
    parser.add_argument(
        "tournament",
        help="USTA Tournament ID or full USTA tournament registration URL."
    )
    parser.add_argument(
        "-d", "--division",
        default=None,
        help="Target division or event name (e.g., 'Boys 14s Singles', 'U12', etc.)."
    )
    parser.add_argument(
        "-o", "--output",
        default="tennislink_report.xlsx",
        help="Output Excel report filename (defaults to 'tennislink_report.xlsx')."
    )
    parser.add_argument(
        "--no-login",
        action="store_true",
        help="Skip UTR login and perform unauthenticated public searches."
    )
    return parser.parse_args()

def format_excel_report(file_path: str):
    """
    Applies custom styling to the generated Excel file to make it visually stunning.
    """
    if not file_path.endswith(".xlsx"):
        return

    logger.info(f"Applying visual styling and formatting to {file_path}...")
    wb = load_workbook(file_path)
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    # Color Palette Constants (Classic Professional Navy Theme)
    HEADER_BG = "1F497D"      # Navy Blue
    HEADER_FG = "FFFFFF"      # White
    ZEBRA_BG = "F2F5F8"       # Light Slate Blue/Gray
    WARN_BG = "FFF2CC"        # Soft Warm Orange/Yellow
    BORDER_COLOR = "D9D9D9"   # Light Gray

    # Styles
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color=HEADER_FG)
    data_font = Font(name=font_family, size=10)
    bold_data_font = Font(name=font_family, size=10, bold=True)
    warn_font = Font(name=font_family, size=10, italic=True)
    
    header_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    zebra_fill = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    warn_fill = PatternFill(start_color=WARN_BG, end_color=WARN_BG, fill_type="solid")
    
    thin_side = Side(border_style="thin", color=BORDER_COLOR)
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    # Header format
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 28

    # Apply data formatting and alternating rows
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_zebra = (row_idx % 2 == 0)
        
        # Check confidence (column index 8 is usually 'Match Confidence')
        confidence_cell = ws.cell(row=row_idx, column=8)
        is_low_confidence = False
        try:
            val = confidence_cell.value
            if val is not None and float(val) < 0.75:
                is_low_confidence = True
        except ValueError:
            pass

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
            # Row Background Fill
            if is_low_confidence:
                cell.fill = warn_fill
                if col_idx == 8:
                    cell.font = warn_font
            elif is_zebra:
                cell.fill = zebra_fill
                
            # Alignments & Number formatting based on column headers
            header_val = ws.cell(row=1, column=col_idx).value or ""
            header_clean = str(header_val).strip().lower()
            
            if any(x in header_clean for x in ["id", "state"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif any(x in header_clean for x in ["wtn", "utr", "confidence", "rank"]):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                # Format numbers
                if "wtn" in header_clean or "utr" in header_clean:
                    cell.number_format = "0.00"
                elif "confidence" in header_clean:
                    cell.number_format = "0.0%"
                elif "rank" in header_clean:
                    cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Autofit column widths with padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                # Format checks for length estimation
                if cell.number_format == "0.0%":
                    val_str = f"{float(val)*100:.1f}%"
                elif cell.number_format == "0.00":
                    val_str = f"{float(val):.2f}"
                else:
                    val_str = str(val)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)
    logger.info("Excel report styled successfully!")

def main():
    args = parse_args()
    
    # 1. Initialize USTA Scraper & Extract Players
    logger.info("Initializing USTA Scraper...")
    usta_scraper = USTAScraper(headless=HEADLESS)
    
    logger.info(f"Scraping USTA for tournament: '{args.tournament}' [Event: {args.division or 'All'}]")
    usta_players = usta_scraper.scrape_tournament(args.tournament, args.division)
    
    if not usta_players:
        logger.error("No players were extracted from the USTA tournament page. Please check the URL/ID or try again.")
        return

    logger.info(f"Successfully retrieved {len(usta_players)} registered players from USTA.")

    # 2. Initialize UTR Scraper & Authenticate
    logger.info("Initializing UTR Scraper...")
    utr_scraper = UTRScraper(headless=HEADLESS)
    
    if args.no_login:
        logger.info("Skipping UTR login as per '--no-login' flag.")
    else:
        logger.info("Logging into UTR...")
        logged_in = utr_scraper.login_and_capture_session()
        if not logged_in:
            logger.warning("Could not authenticate with UTR. Search heuristics will proceed in public mode.")

    # 3. Initialize Matcher
    matcher = PlayerMatcher(utr_scraper)

    # 4. Resolve player mappings and ratings
    resolved_records = []
    
    logger.info("Resolving player UTR ratings and mapping profiles...")
    for idx, up in enumerate(usta_players):
        logger.info(f"[{idx+1}/{len(usta_players)}] Resolving player: {up['name']}")
        
        # Match using the heuristics engine
        match_res = matcher.find_utr_profile(up, usta_players)
        
        if match_res:
            record = {
                "Player Name": up["name"],
                "Hometown": f"{up.get('city') or ''}, {up.get('state') or ''}".strip(", "),
                "USTA ID": up["usta_id"],
                "USTA Ranking": up.get("ranking") or "N/A",
                "WTN Singles": up.get("wtn_singles") or "N/A",
                "WTN Doubles": up.get("wtn_doubles") or "N/A",
                "UTR ID": match_res["utr_id"],
                "Match Confidence": match_res["confidence"],
                "Match Method": match_res["match_method"],
                "UTR Singles": match_res.get("utr_singles") or "N/A",
                "UTR Doubles": match_res.get("utr_doubles") or "N/A"
            }
        else:
            record = {
                "Player Name": up["name"],
                "Hometown": f"{up.get('city') or ''}, {up.get('state') or ''}".strip(", "),
                "USTA ID": up["usta_id"],
                "USTA Ranking": up.get("ranking") or "N/A",
                "WTN Singles": up.get("wtn_singles") or "N/A",
                "WTN Doubles": up.get("wtn_doubles") or "N/A",
                "UTR ID": "N/A",
                "Match Confidence": 0.0,
                "Match Method": "None",
                "UTR Singles": "N/A",
                "UTR Doubles": "N/A"
            }
            
        resolved_records.append(record)

    # 5. Export to Excel
    df = pd.DataFrame(resolved_records)
    
    # Reorder columns to make the sheet look elegant
    col_order = [
        "Player Name", "Hometown", "USTA ID", "USTA Ranking", "WTN Singles", "WTN Doubles",
        "UTR ID", "Match Confidence", "UTR Singles", "UTR Doubles", "Match Method"
    ]
    df = df[col_order]

    output_file = args.output
    if not output_file.endswith(".xlsx") and not output_file.endswith(".csv"):
        output_file += ".xlsx"

    logger.info(f"Saving compiled data to {output_file}...")
    if output_file.endswith(".csv"):
        df.to_csv(output_file, index=False)
    else:
        df.to_excel(output_file, index=False)
        # Apply premium formatting
        format_excel_report(output_file)

    logger.info(f"Done! All registered players processed. Report generated successfully at: {output_file}")

if __name__ == "__main__":
    main()
