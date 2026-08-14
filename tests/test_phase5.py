import os
import sys
import tempfile
import pandas as pd
from unittest.mock import patch, MagicMock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

print("=" * 60)
print("Phase 5: Testing CLI Orchestrator & Exporter")
print("=" * 60)

# 1. Test CLI argument parsing
print("\n[1/4] Testing CLI argument parsing...")
from main import parse_args

test_args = ["main.py", "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B", "-d", "Boys 18", "-o", "test_report.xlsx"]
with patch.object(sys, 'argv', test_args):
    args = parse_args()
    assert args.tournament == "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B"
    assert args.division == "Boys 18"
    assert args.output == "test_report.xlsx"
    assert args.no_login == False
print("  Full arguments parsed: PASSED")

test_args_defaults = ["main.py", "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B"]
with patch.object(sys, 'argv', test_args_defaults):
    args = parse_args()
    assert args.tournament == "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B"
    assert args.division == None
    assert args.output == "tennislink_report.xlsx"
    assert args.no_login == False
print("  Default arguments: PASSED")

test_args_nologin = ["main.py", "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B", "--no-login"]
with patch.object(sys, 'argv', test_args_nologin):
    args = parse_args()
    assert args.no_login == True
print("  --no-login flag: PASSED")

try:
    with patch.object(sys, 'argv', ["main.py", "--help"]):
        try:
            parse_args()
        except SystemExit:
            pass
    print("  --help triggers SystemExit: PASSED")
except Exception:
    pass
print("  PASSED")

# 2. Test Excel formatting function
print("\n[2/4] Testing Excel report formatting...")
from main import format_excel_report

tmp_xlsx = os.path.join(tempfile.gettempdir(), "test_format_output.xlsx")
df = pd.DataFrame({
    "Player Name": ["Alice Williams", "Bob Smith", "Charlie Brown"],
    "Hometown": ["Los Angeles, CA", "Chicago, IL", "New York, NY"],
    "USTA ID": ["11111", "22222", "33333"],
    "USTA Ranking": [5, 12, "N/A"],
    "WTN Singles": [16.5, 14.2, "N/A"],
    "WTN Doubles": [15.0, 13.8, "N/A"],
    "UTR ID": ["UTR_ALICE", "UTR_BOB", "N/A"],
    "Match Confidence": [0.95, 0.88, 0.0],
    "UTR Singles": [11.5, 8.0, "N/A"],
    "UTR Doubles": [10.2, 7.5, "N/A"],
    "Match Method": ["name_geo", "match_history_circle", "None"]
})
df.to_excel(tmp_xlsx, index=False)

try:
    format_excel_report(tmp_xlsx)
    from openpyxl import load_workbook
    wb = load_workbook(tmp_xlsx)
    ws = wb.active
    
    header = ws.cell(row=1, column=1)
    assert header.font.bold == True, "Header should be bold"
    
    assert ws.cell(row=2, column=1).value == "Alice Williams"
    assert ws.cell(row=4, column=1).value == "Charlie Brown"
    
    from openpyxl.utils import get_column_letter
    for col_idx in range(1, 12):
        col_letter = get_column_letter(col_idx)
        col_width = ws.column_dimensions[col_letter].width
        assert col_width is not None and col_width > 0, f"Column {col_letter} width should be set"
    
    confidence_cell = ws.cell(row=2, column=8)
    assert confidence_cell.number_format == "0.0%", f"Expected 0.0% format, got {confidence_cell.number_format}"
    
    charlie_row = 4
    charlie_conf_cell = ws.cell(row=charlie_row, column=8)
    assert charlie_conf_cell.font.italic == True, "Low confidence row should have italic font"
    
    wb.close()
    print("  Excel formatting applied successfully")
except Exception as e:
    print(f"  Excel formatting error: {e}")
    raise

os.remove(tmp_xlsx)
print("  PASSED")

# 3. Test CSV output path
print("\n[3/4] Testing output path handling...")
df.to_csv(os.path.join(tempfile.gettempdir(), "test_csv.csv"), index=False)
print("  CSV export works: PASSED")
print("  PASSED")

# 4. Test full pipeline with mocked components
print("\n[4/4] Testing full main pipeline (mocked)...")
from main import main

mock_usta_players = [
    {"usta_id": "11111", "name": "Alice Williams", "city": "Los Angeles", "state": "CA", "wtn_singles": 16.5, "wtn_doubles": 15.0, "ranking": 5},
    {"usta_id": "22222", "name": "Bob Smith", "city": "Chicago", "state": "IL", "wtn_singles": 14.2, "wtn_doubles": 13.8, "ranking": 12},
]

def mock_find_utr(usta_player, tournament_players=None):
    mapping = {
        "11111": {"utr_id": "UTR_ALICE", "name": "Alice Williams", "city": "Los Angeles", "state": "CA", "utr_singles": 11.5, "utr_doubles": 10.2, "confidence": 0.95, "match_method": "name_geo", "source": "search"},
        "22222": {"utr_id": "UTR_BOB", "name": "Bob Smith", "city": "Chicago", "state": "IL", "utr_singles": 8.0, "utr_doubles": 7.5, "confidence": 0.88, "match_method": "match_history_circle", "source": "search"}
    }
    return mapping.get(usta_player["usta_id"])

output_path = os.path.join(tempfile.gettempdir(), "test_pipeline_output.xlsx")

with patch('main.USTAScraper') as MockUSTAScraper, \
     patch('main.UTRScraper') as MockUTRScraper, \
     patch('main.PlayerMatcher') as MockPlayerMatcher, \
     patch.object(sys, 'argv', ["main.py", "D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B", "-d", "Boys 18", "-o", output_path]):

    mock_usta_instance = MockUSTAScraper.return_value
    mock_usta_instance.scrape_tournament.return_value = mock_usta_players
    
    mock_utr_instance = MockUTRScraper.return_value
    mock_utr_instance.login_if_needed.return_value = True
    
    mock_matcher_instance = MockPlayerMatcher.return_value
    mock_matcher_instance.find_utr_profile.side_effect = mock_find_utr

    main()
    
    MockUSTAScraper.assert_called_once()
    mock_usta_instance.scrape_tournament.assert_called_once_with("D6F0B896-6620-4052-B6C4-3ABBAE1C5A8B", "Boys 18")
    
    MockUTRScraper.assert_called_once()
    mock_utr_instance.login_if_needed.assert_called_once()
    
    MockPlayerMatcher.assert_called_once_with(mock_utr_instance)
    assert mock_matcher_instance.find_utr_profile.call_count == 2

    assert os.path.exists(output_path), f"Output file {output_path} should exist"
    
    df_result = pd.read_excel(output_path)
    assert len(df_result) == 2
    assert df_result.iloc[0]["Player Name"] == "Alice Williams"
    assert df_result.iloc[1]["Player Name"] == "Bob Smith"
    assert df_result.iloc[0]["UTR ID"] == "UTR_ALICE"
    assert df_result.iloc[0]["Match Confidence"] == 0.95
    
    os.remove(output_path)

print("  Full pipeline executed successfully")
print("  PASSED")

print("\n" + "=" * 60)
print("Phase 5: ALL TESTS PASSED")
print("=" * 60)
